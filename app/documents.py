import io
from html import escape
from decimal import Decimal
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import Response
from PIL import Image as PILImage, UnidentifiedImageError
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
from docx import Document as WordDocument
from docx.shared import Inches, RGBColor
from sqlalchemy import select, func
from sqlalchemy.orm import Session
from .core.database import get_db
from .core.security import current_user, require, audit
from .core.models import uid
from .models import (
    Document,
    School,
    Student,
    Term,
    AcademicSession,
    ClassArm,
    ClassLevel,
    Result,
    ResultScore,
    AssessmentComponent,
    Subject,
    Attendance,
    Payment,
    Parent,
    StudentParent,
    User,
    GradingScale,
    GradingScaleItem,
)
from .resources import scoped, class_access
from .storage import Storage

router = APIRouter(tags=["Documents and branding"])
MIMES = {"pdf": "application/pdf", "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}


def document_content(db, document):
    school = db.get(School, document.school_id)
    student = db.get(Student, document.student_id)
    term = db.get(Term, document.term_id)
    session = db.get(AcademicSession, term.session_id)
    class_arm = db.get(ClassArm, student.class_id) if student.class_id else None
    class_label = (db.get(ClassLevel, class_arm.level_id).name + " " + class_arm.name) if class_arm else "Unassigned"
    intro = [
        f"{student.first_name} {student.middle_name} {student.last_name}".replace("  ", " "),
        f"Student ID: {student.student_code}",
        f"Class: {class_label}",
        f"Session: {session.name} | Term: {term.name}",
    ]
    if document.kind == "RECEIPT":
        payment = db.get(Payment, document.payment_id)
        parents = db.scalars(
            select(Parent.name)
            .join(StudentParent, StudentParent.parent_id == Parent.id)
            .where(StudentParent.student_id == student.id, Parent.school_id == school.id)
        ).all()
        intro += ["Parent / guardian: " + ", ".join(parents), "Receipt: " + payment.receipt_number]
        rows = [
            ["Payment details", "Value"],
            ["Date", str(payment.payment_date)],
            ["Method", payment.method],
            ["Reference", payment.reference or "—"],
            ["Previous balance", f"{school.currency} {payment.previous_balance:,.2f}"],
            ["Amount received", f"{school.currency} {payment.amount:,.2f}"],
            ["New balance", f"{school.currency} {payment.new_balance:,.2f}"],
            ["Recorded by", db.get(User, payment.recorded_by).name],
        ]
        footer = (
            ["REVERSED: " + payment.reversal_reason]
            if payment.reversed
            else ["Thank you. Keep this receipt for your records."]
        )
        return school, "PAYMENT RECEIPT", intro, rows, footer
    results = db.scalars(
        select(Result)
        .where(Result.school_id == school.id, Result.student_id == student.id, Result.term_id == term.id)
        .order_by(Result.subject_id)
    ).all()
    if not results or any(r.status != "PUBLISHED" for r in results):
        raise ValueError("All results must be published before generating a report card.")
    rows = [["Subject", "Assessments", "Total", "Grade", "Remark"]]
    for result in results:
        scores = db.scalars(select(ResultScore).where(ResultScore.result_id == result.id)).all()
        components = "; ".join(
            f"{db.get(AssessmentComponent, s.component_id).name}: {s.score:g}/{s.max_score:g}" for s in scores
        )
        rows.append(
            [db.get(Subject, result.subject_id).name, components, str(result.total), result.grade, result.remark]
        )
    attendance = dict(
        db.execute(
            select(Attendance.status, func.count())
            .where(
                Attendance.school_id == school.id, Attendance.student_id == student.id, Attendance.term_id == term.id
            )
            .group_by(Attendance.status)
        ).all()
    )
    total = sum((r.total for r in results), Decimal(0))
    footer = [
        f"Total: {total:.2f} | Average: {total / len(results):.2f}%",
        "Attendance: " + ", ".join(f"{status.title()}: {count}" for status, count in attendance.items()),
    ]
    scale = db.scalar(select(GradingScale).where(GradingScale.school_id == school.id, GradingScale.active.is_(True)))
    if scale:
        grades = db.scalars(
            select(GradingScaleItem)
            .where(GradingScaleItem.scale_id == scale.id)
            .order_by(GradingScaleItem.min_score.desc())
        ).all()
        footer.append("Grading: " + " | ".join(f"{g.label}: {g.min_score:g}–{g.max_score:g}" for g in grades))
    footer += ["Class teacher signature: __________________", "Principal signature: __________________"]
    return school, "STUDENT REPORT CARD", intro, rows, footer


def generate(db, document):
    school, title, intro, rows, footer = document_content(db, document)
    storage, output = Storage(), io.BytesIO()
    logo = storage.get(school.logo_key) if school.logo_key else None
    if document.format == "docx":
        word = WordDocument()
        if logo:
            word.add_picture(io.BytesIO(logo), width=Inches(0.65))
        heading = word.add_heading(school.name, 0)
        for run in heading.runs:
            run.font.color.rgb = RGBColor.from_string(school.primary_color.lstrip("#"))
        word.add_paragraph(f"{school.address}\n{school.email} · {school.phone}")
        word.add_heading(title, 1)
        for line in intro:
            word.add_paragraph(line)
        table = word.add_table(rows=0, cols=len(rows[0]))
        table.style = "Light Shading Accent 1"
        for row in rows:
            for cell, value in zip(table.add_row().cells, row):
                cell.text = str(value)
        for line in footer:
            word.add_paragraph(line)
        word.save(output)
    else:
        styles = getSampleStyleSheet()
        styles["Title"].textColor = colors.HexColor(school.primary_color)
        styles["BodyText"].fontSize = 9
        story = []
        if logo:
            story.append(Image(io.BytesIO(logo), width=48, height=48, kind="proportional"))
        story += [
            Paragraph(escape(school.name), styles["Title"]),
            Paragraph(escape(f"{school.address} | {school.email} | {school.phone}"), styles["BodyText"]),
            Spacer(1, 18),
            Paragraph(title, styles["Heading2"]),
        ]
        story += [Paragraph(escape(line), styles["BodyText"]) for line in intro]
        story.append(Spacer(1, 16))
        table = Table(
            [[Paragraph(escape(str(value)), styles["BodyText"]) for value in row] for row in rows],
            repeatRows=1,
            hAlign="LEFT",
            colWidths=([105, 160, 45, 45, 140] if len(rows[0]) == 5 else [200, 295]),
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e5eee9")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f6f8f7")]),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), 9),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
                    ("LINEBELOW", (0, 0), (-1, 0), 1, colors.HexColor(school.primary_color)),
                ]
            )
        )
        story += [table, Spacer(1, 18)] + [Paragraph(escape(line), styles["BodyText"]) for line in footer]
        SimpleDocTemplate(
            output,
            pagesize=A4,
            rightMargin=40,
            leftMargin=40,
            topMargin=36,
            bottomMargin=36,
            title=title,
            author=school.name,
        ).build(story)
    key = f"{school.id}/documents/{document.id}.{document.format}"
    storage.put(key, output.getvalue(), MIMES[document.format])
    return key


def assert_document_access(db, user, document):
    if user.role == "PARENT":
        linked = db.scalar(
            select(StudentParent.id)
            .join(Parent, Parent.id == StudentParent.parent_id)
            .where(
                Parent.school_id == user.school_id,
                Parent.user_id == user.id,
                StudentParent.student_id == document.student_id,
            )
        )
        if not linked:
            raise HTTPException(404, "Document was not found.")
    else:
        require(user, "finance.view" if document.kind == "RECEIPT" else "results.view")
        student = scoped(db, Student, document.student_id, user)
        class_access(db, user, student.class_id)
        from .resources import class_teacher_ids

        if user.role == "TEACHER" and student.class_id not in class_teacher_ids(db, user):
            raise HTTPException(403, "Report cards require a class-teacher assignment.")


@router.get("/documents/{document_id}/download")
def download(document_id: str, user=Depends(current_user), db: Session = Depends(get_db, scope="function")):
    document = scoped(db, Document, document_id, user)
    assert_document_access(db, user, document)
    if document.kind == "RECEIPT" and db.get(Payment, document.payment_id).reversed:
        raise HTTPException(409, "This payment was reversed. The original receipt is no longer valid.")
    if document.status != "READY" or not document.storage_key:
        raise HTTPException(409, "Document is not ready or has been superseded.")
    return Response(
        Storage().get(document.storage_key),
        media_type=MIMES[document.format],
        headers={
            "Content-Disposition": f'attachment; filename="ile-iwe-{document.kind.lower()}-{document.id[:8]}.{document.format}"',
            "Cache-Control": "private, no-store",
        },
    )


@router.post("/school/logo")
def upload_logo(
    file: UploadFile = File(...), user=Depends(current_user), db: Session = Depends(get_db, scope="function")
):
    require(user, "school.manage_branding")
    content = file.file.read(2 * 1024 * 1024 + 1)
    if len(content) > 2 * 1024 * 1024:
        raise HTTPException(422, "Logo must be smaller than 2 MB.")
    try:
        image = PILImage.open(io.BytesIO(content))
        if image.width * image.height > 16_000_000:
            raise ValueError()
        image.thumbnail((600, 600))
        safe = io.BytesIO()
        image.convert("RGBA").save(safe, "PNG")
    except (UnidentifiedImageError, OSError, ValueError, PILImage.DecompressionBombError):
        raise HTTPException(422, "Upload a valid PNG, JPEG, or WebP image.")
    key = f"{user.school_id}/branding/{uid()}.png"
    Storage().put(key, safe.getvalue(), "image/png")
    school = db.get(School, user.school_id)
    school.logo_key = key
    audit(db, user, "school.logo_updated", "schools", school.id)
    return {"has_logo": True}


@router.get("/school/logo")
def logo(user=Depends(current_user), db: Session = Depends(get_db, scope="function")):
    school = db.get(School, user.school_id)
    if not school.logo_key:
        raise HTTPException(404, "No logo has been uploaded.")
    return Response(
        Storage().get(school.logo_key), media_type="image/png", headers={"Cache-Control": "private, max-age=300"}
    )


@router.post("/students/{student_id}/photo")
def upload_student_photo(
    student_id: str,
    file: UploadFile = File(...),
    user=Depends(current_user),
    db: Session = Depends(get_db, scope="function"),
):
    require(user, "students.update")
    student = scoped(db, Student, student_id, user, lock=True)
    content = file.file.read(2 * 1024 * 1024 + 1)
    if len(content) > 2 * 1024 * 1024:
        raise HTTPException(422, "Photo must be smaller than 2 MB.")
    try:
        image = PILImage.open(io.BytesIO(content))
        if image.width * image.height > 16_000_000:
            raise ValueError()
        image.thumbnail((800, 800))
        output = io.BytesIO()
        image.convert("RGB").save(output, "JPEG", quality=88)
    except (UnidentifiedImageError, OSError, ValueError, PILImage.DecompressionBombError):
        raise HTTPException(422, "Upload a valid PNG, JPEG, or WebP photo.")
    key = f"{user.school_id}/students/{student.id}/{uid()}.jpg"
    Storage().put(key, output.getvalue(), "image/jpeg")
    student.photo_key = key
    audit(db, user, "student.photo_updated", "students", student.id)
    return {"has_photo": True}


@router.get("/students/{student_id}/photo")
def student_photo(student_id: str, user=Depends(current_user), db: Session = Depends(get_db, scope="function")):
    if user.role == "PARENT":
        from .resources import child_access

        student = child_access(db, user, student_id)
    else:
        require(user, "students.view")
        student = scoped(db, Student, student_id, user)
        class_access(db, user, student.class_id)
    if not student.photo_key:
        raise HTTPException(404, "No student photo has been uploaded.")
    return Response(
        Storage().get(student.photo_key), media_type="image/jpeg", headers={"Cache-Control": "private, no-store"}
    )


@router.post("/staff/{staff_id}/photo")
def upload_staff_photo(
    staff_id: str,
    file: UploadFile = File(...),
    user=Depends(current_user),
    db: Session = Depends(get_db, scope="function"),
):
    from .models import Staff

    staff = scoped(db, Staff, staff_id, user, lock=True)
    if staff.user_id != user.id:
        require(user, "staff.manage")
    content = file.file.read(2 * 1024 * 1024 + 1)
    if len(content) > 2 * 1024 * 1024:
        raise HTTPException(422, "Photo must be smaller than 2 MB.")
    try:
        image = PILImage.open(io.BytesIO(content))
        if image.width * image.height > 16_000_000:
            raise ValueError()
        image.thumbnail((800, 800))
        output = io.BytesIO()
        image.convert("RGB").save(output, "JPEG", quality=88)
    except (UnidentifiedImageError, OSError, ValueError, PILImage.DecompressionBombError):
        raise HTTPException(422, "Upload a valid PNG, JPEG, or WebP photo.")
    staff.photo_key = f"{user.school_id}/staff/{staff.id}/{uid()}.jpg"
    Storage().put(staff.photo_key, output.getvalue(), "image/jpeg")
    audit(db, user, "staff.photo_updated", "staff", staff.id)
    return {"has_photo": True}


@router.get("/staff/{staff_id}/photo")
def staff_photo(staff_id: str, user=Depends(current_user), db: Session = Depends(get_db, scope="function")):
    from .models import Staff

    staff = scoped(db, Staff, staff_id, user)
    if staff.user_id != user.id:
        require(user, "staff.view")
    if not staff.photo_key:
        raise HTTPException(404, "No photo has been uploaded.")
    return Response(
        Storage().get(staff.photo_key), media_type="image/jpeg", headers={"Cache-Control": "private, no-store"}
    )
